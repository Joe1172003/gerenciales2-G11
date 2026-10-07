"""
Persona 5 - Genera los archivos de importacion de Empleados para Odoo 18.

Salida (carpeta empleados/ en la raiz de Proyecto1):
  1_departamentos.xlsx  -> Empleados > Configuracion > Departamentos > Importar
  2_cargos.xlsx         -> Empleados > Configuracion > Puestos de trabajo > Importar
  3_empleados.xlsx      -> Empleados > Empleados > Importar
  4_jefes_departamento.xlsx -> Departamentos > Importar (otra vez, solo actualiza el jefe)

Se importan EN ESE ORDEN. Los encabezados usan nombres tecnicos, asi el
importador los reconoce solo sin importar el idioma del usuario.

La columna "id" es el ID externo: si se vuelve a importar el mismo archivo,
Odoo ACTUALIZA los registros en vez de duplicarlos.

Uso:  python scripts/p5/generar_empleados.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

RAIZ = Path(__file__).resolve().parents[2]
SALIDA = RAIZ / "empleados"

DOMINIO = "quetzalmart.example.com"  # correos de prueba, nunca reales

# --- 5 departamentos ---------------------------------------------------------
DEPARTAMENTOS = [
    ("qm_dep_gerencia", "Gerencia General"),
    ("qm_dep_ventas", "Ventas y Atención al Cliente"),
    ("qm_dep_compras", "Compras y Logística"),
    ("qm_dep_bodega", "Bodega e Inventario"),
    ("qm_dep_rrhh", "Recursos Humanos"),
]

# --- 6 cargos (puesto -> departamento) ---------------------------------------
CARGOS = [
    ("qm_job_gerente", "Gerente de Sucursal", "qm_dep_gerencia"),
    ("qm_job_cajero", "Cajero", "qm_dep_ventas"),
    ("qm_job_asesor", "Asesor de Ventas", "qm_dep_ventas"),
    ("qm_job_compras", "Encargado de Compras", "qm_dep_compras"),
    ("qm_job_bodeguero", "Bodeguero", "qm_dep_bodega"),
    ("qm_job_rrhh", "Analista de Recursos Humanos", "qm_dep_rrhh"),
]
DEP_DE_CARGO = {c[0]: c[2] for c in CARGOS}

SUCURSALES = {
    "GT": ("QM Guatemala", "+502"),
    "MX": ("QM México", "+52"),
    "SV": ("QM El Salvador", "+503"),
}

# --- 35 empleados: (nombre, genero, sucursal, cargo) -------------------------
# Guatemala (central) 15 · Mexico 10 · El Salvador 10
EMPLEADOS = [
    # Guatemala
    ("Ana Lucía Morales Pérez", "female", "GT", "qm_job_gerente"),
    ("Carlos Eduardo Ajú Tzul", "male", "GT", "qm_job_compras"),
    ("María Fernanda López Ixcoy", "female", "GT", "qm_job_rrhh"),
    ("José Antonio Cifuentes Ramírez", "male", "GT", "qm_job_bodeguero"),
    ("Sofía Alejandra Herrera Batz", "female", "GT", "qm_job_cajero"),
    ("Luis Fernando Gómez Chávez", "male", "GT", "qm_job_asesor"),
    ("Andrea Beatriz Castillo Ortiz", "female", "GT", "qm_job_cajero"),
    ("Diego Alejandro Monterroso Paz", "male", "GT", "qm_job_bodeguero"),
    ("Gabriela Isabel Xicará Méndez", "female", "GT", "qm_job_asesor"),
    ("Kevin Rodrigo Estrada Juárez", "male", "GT", "qm_job_compras"),
    ("Daniela Marisol Coyoy García", "female", "GT", "qm_job_cajero"),
    ("Javier Ernesto Recinos Arana", "male", "GT", "qm_job_asesor"),
    ("Paola Verónica Sian Hernández", "female", "GT", "qm_job_rrhh"),
    ("Óscar Daniel Pineda Solares", "male", "GT", "qm_job_bodeguero"),
    ("Mónica Raquel Tum Alvarado", "female", "GT", "qm_job_cajero"),
    # Mexico
    ("Alejandro Sánchez Villarreal", "male", "MX", "qm_job_gerente"),
    ("Valeria Guadalupe Torres Ríos", "female", "MX", "qm_job_compras"),
    ("Miguel Ángel Ramírez Fuentes", "male", "MX", "qm_job_bodeguero"),
    ("Ximena Hernández Zúñiga", "female", "MX", "qm_job_cajero"),
    ("Jorge Luis Domínguez Cano", "male", "MX", "qm_job_asesor"),
    ("Fernanda Ruiz Cervantes", "female", "MX", "qm_job_cajero"),
    ("Ricardo Mendoza Ochoa", "male", "MX", "qm_job_bodeguero"),
    ("Itzel Martínez Salgado", "female", "MX", "qm_job_asesor"),
    ("Emiliano Vargas Lozano", "male", "MX", "qm_job_rrhh"),
    ("Regina Navarro Guzmán", "female", "MX", "qm_job_cajero"),
    # El Salvador
    ("Roberto Carlos Hernández Flores", "male", "SV", "qm_job_gerente"),
    ("Karla Patricia Martínez Rivas", "female", "SV", "qm_job_compras"),
    ("Wilfredo Antonio Guardado Mejía", "male", "SV", "qm_job_bodeguero"),
    ("Claudia Elizabeth Portillo Ayala", "female", "SV", "qm_job_cajero"),
    ("Mauricio Ernesto Alvarenga Cruz", "male", "SV", "qm_job_asesor"),
    ("Yesenia Marlene Quintanilla Romero", "female", "SV", "qm_job_cajero"),
    ("Héctor Manuel Rivera Orellana", "male", "SV", "qm_job_bodeguero"),
    ("Jennifer Alexandra Bonilla Castro", "female", "SV", "qm_job_asesor"),
    ("Francisco Javier Menjívar Pleitez", "male", "SV", "qm_job_rrhh"),
    ("Silvia Carolina Escobar Molina", "female", "SV", "qm_job_cajero"),
]

# Jefe de cada departamento (codigo de empleado). Al importarlo, Odoo pone a
# esa persona como "Gerente" de todos los empleados del departamento.
JEFES = {
    "qm_dep_gerencia": "EMP001",  # Ana Lucía Morales (Gerente QM Guatemala, central)
    "qm_dep_ventas": "EMP006",    # Luis Fernando Gómez
    "qm_dep_compras": "EMP002",   # Carlos Eduardo Ajú
    "qm_dep_bodega": "EMP004",    # José Antonio Cifuentes
    "qm_dep_rrhh": "EMP003",      # María Fernanda López
}

# Empleados con contrato en el gestor documental (ver generar_contratos.py)
CON_CONTRATO = ["EMP001", "EMP005", "EMP016", "EMP019", "EMP026"]


def _usuario_correo(nombre):
    import unicodedata
    limpio = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode().lower()
    partes = limpio.split()
    apellido = partes[-2] if len(partes) >= 3 else partes[-1]
    return f"{partes[0]}.{apellido}"


def empleados():
    """Lista de dicts con todos los datos de cada empleado (la usa tambien
    generar_contratos.py, asi los contratos llevan los mismos nombres)."""
    salida, usados = [], set()
    for i, (nombre, genero, suc, cargo) in enumerate(EMPLEADOS, start=1):
        codigo = f"EMP{i:03d}"
        usuario = _usuario_correo(nombre)
        if usuario in usados:
            usuario = f"{usuario}{i}"
        usados.add(usuario)
        sucursal, lada = SUCURSALES[suc]
        salida.append({
            "codigo": codigo,
            "id": f"qm_{codigo.lower()}",
            "nombre": nombre,
            "genero": genero,
            "pais": suc,
            "sucursal": sucursal,
            "cargo_id": cargo,
            "cargo": next(c[1] for c in CARGOS if c[0] == cargo),
            "departamento_id": DEP_DE_CARGO[cargo],
            "departamento": next(d[1] for d in DEPARTAMENTOS if d[0] == DEP_DE_CARGO[cargo]),
            "correo": f"{usuario}@{DOMINIO}",
            "telefono": f"{lada} {5000 + i * 37:04d}-{1000 + i * 113 % 9000:04d}",
            "identificacion": f"{suc}-{2600000 + i * 7919}",
        })
    return salida


def _hoja(ruta, encabezados, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Importar"
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for f in filas:
        ws.append(f)
    for col in ws.columns:
        ancho = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(ancho + 2, 45)
    ws.freeze_panes = "A2"
    wb.save(ruta)


def main():
    SALIDA.mkdir(exist_ok=True)
    emps = empleados()
    por_codigo = {e["codigo"]: e for e in emps}

    _hoja(SALIDA / "1_departamentos.xlsx", ["id", "name"],
          [list(d) for d in DEPARTAMENTOS])

    _hoja(SALIDA / "2_cargos.xlsx", ["id", "name", "department_id/id"],
          [list(c) for c in CARGOS])

    _hoja(SALIDA / "3_empleados.xlsx",
          ["id", "name", "barcode", "job_id/id", "job_title", "department_id/id",
           "category_ids", "work_email", "mobile_phone", "identification_id", "gender"],
          [[e["id"], e["nombre"], e["codigo"], e["cargo_id"], e["cargo"],
            e["departamento_id"], f"{e['sucursal']},Empleado", e["correo"],
            e["telefono"], e["identificacion"], e["genero"]] for e in emps])

    _hoja(SALIDA / "4_jefes_departamento.xlsx", ["id", "manager_id/id"],
          [[dep, por_codigo[cod]["id"]] for dep, cod in JEFES.items()])

    print(f"Listo: {len(DEPARTAMENTOS)} departamentos, {len(CARGOS)} cargos, "
          f"{len(emps)} empleados -> {SALIDA}")


if __name__ == "__main__":
    main()
