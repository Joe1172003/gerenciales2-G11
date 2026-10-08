"""
Persona 5 - Version de REFERENCIA de lo que hace el bot de UiPath con los Excel.

NO es el bot. El bot es UiPath (ver GUIA_BOT_UIPATH.md). Este script hace la
misma transformacion que LimpiarClientes.vb y LimpiarProductos.vb, para:
  1. Generar la "salida esperada" y probar la importacion en Odoo a mano
     antes de tener el bot listo.
  2. Comparar: si el bot y este script producen lo mismo, el bot esta bien.

Uso:
  python scripts/p5/rpa/referencia_bot.py <carpeta_entrada> <carpeta_salida> [ubicacion] [--no-publicar]
  ej: python scripts/p5/rpa/referencia_bot.py carpeta_prueba_rpa salida_esperada "GT/Existencias" --no-publicar
"""
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook

COLS_CLIENTES = ["id", "name", "company_type", "parent_id", "email", "phone", "street", "street2", "city",
                 "state_id", "zip", "country_id", "vat", "Website Link", "category_id/id", "ref", "comment"]
COLS_ETIQUETAS = ["id", "name"]
COLS_PRODUCTOS = ["id", "name", "type", "default_code", "barcode", "list_price", "standard_price", "weight",
                  "description_sale", "is_storable", "is_published"]
COLS_EXISTENCIAS = ["product_id", "location_id", "inventory_quantity"]

TIPOS_ALMACENABLES = {"goods", "consu", "bienes", "consumible", "producto", "storable", "almacenable"}
VERDADEROS = {"true", "1", "yes", "si", "verdadero", "x"}


def norm(s):
    t = (s or "").replace("*", " ").strip().lower()
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t).strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", norm(s)).strip("_")


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if abs(v - round(v)) < 0.0001:
            return str(int(round(v)))
        return f"{v:.12f}".rstrip("0").rstrip(".")
    s = str(v).strip()
    if re.fullmatch(r"-?\d+\.\d+", s) and abs(float(s) - round(float(s))) < 0.0001:
        return str(int(round(float(s))))
    return s


def valor(fila, mapa, claves):
    for k in claves.split("|"):
        if k in mapa:
            t = texto(fila[mapa[k]])
            if t:
                return t
    return ""


def limpiar_clientes(hoja, clientes, etiquetas):
    filas = list(hoja.iter_rows(values_only=True))
    mapa = {}
    for i, h in enumerate(filas[0] if filas else []):
        mapa.setdefault(norm(str(h or "")), i)
    agregados = omitidos = 0
    ids = {c[0] for c in clientes}
    for f in filas[1:]:
        nombre = valor(f, mapa, "name|nombre")
        if not nombre:
            continue
        ref, email, vat = valor(f, mapa, "reference|referencia"), valor(f, mapa, "email|correo"), valor(f, mapa, "tax id|nit")
        xid = "rpa_cli_" + slug(ref or vat or email or nombre)
        if xid in ids:
            omitidos += 1
            continue
        tags = ["rpa_tag_rpa"]
        for t in valor(f, mapa, "tags|etiquetas").split(","):
            t = t.strip()
            if not t:
                continue
            tid = "rpa_tag_" + slug(t)
            if tid not in tags:
                tags.append(tid)
            if tid not in {e[0] for e in etiquetas}:
                etiquetas.append([tid, t])
        clientes.append([xid, nombre, valor(f, mapa, "company type|tipo de compania"),
                         valor(f, mapa, "related company|compania relacionada"), email,
                         valor(f, mapa, "phone|telefono"), valor(f, mapa, "street|calle"),
                         valor(f, mapa, "street2|calle 2"), valor(f, mapa, "city|ciudad"),
                         valor(f, mapa, "state|estado|departamento"), valor(f, mapa, "zip|codigo postal"),
                         valor(f, mapa, "country|pais"), vat, valor(f, mapa, "website|sitio web"),
                         ",".join(tags), ref, valor(f, mapa, "notes|notas")])
        ids.add(xid)
        agregados += 1
    return agregados, omitidos


def limpiar_productos(hoja, productos, existencias, ubicacion, no_publicar):
    filas = list(hoja.iter_rows(values_only=True))
    mapa = {}
    for i, h in enumerate(filas[0] if filas else []):
        mapa.setdefault(norm(str(h or "")), i)
    agregados = omitidos = 0
    ids = {p[0] for p in productos}
    for f in filas[1:]:
        nombre = valor(f, mapa, "name|nombre")
        if not nombre:
            continue
        xid = valor(f, mapa, "external id|id externo|id") or ("rpa_prd_" + slug(nombre))
        if xid in ids:
            omitidos += 1
            continue
        tipo = valor(f, mapa, "product type|tipo de producto") or "Goods"
        almacenable = norm(tipo) in TIPOS_ALMACENABLES
        publicado = (not no_publicar) and norm(valor(f, mapa, "esta publicado|is published|published")) in VERDADEROS
        codigo = valor(f, mapa, "internal reference|referencia interna")
        productos.append([xid, nombre, tipo, codigo, valor(f, mapa, "barcode|codigo de barras"),
                          valor(f, mapa, "sales price|precio de venta"), valor(f, mapa, "cost|costo"),
                          valor(f, mapa, "weight|peso"), valor(f, mapa, "sales description|descripcion de venta"),
                          "True" if almacenable else "False", "True" if publicado else "False"])
        ids.add(xid)
        agregados += 1
        cantidad = valor(f, mapa, "cantidad a la mano|quantity on hand|on hand")
        if almacenable and cantidad not in ("", "0"):
            existencias.append([codigo or nombre, ubicacion, cantidad])
    return agregados, omitidos


def escribir(ruta, columnas, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(columnas)
    for f in filas:
        ws.append(f)
    wb.save(ruta)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    entrada, salida = Path(args[0]), Path(args[1])
    ubicacion = args[2] if len(args) > 2 else "QMGT/Existencias"
    no_publicar = "--no-publicar" in sys.argv
    salida.mkdir(parents=True, exist_ok=True)

    clientes, etiquetas, productos, existencias = [], [["rpa_tag_rpa", "RPA"]], [], []
    for archivo in sorted(list(entrada.rglob("*.xlsx")) + list(entrada.rglob("*.xls"))):
        if archivo.name.startswith("~$"):
            continue
        wb = load_workbook(archivo, read_only=True, data_only=True)
        for nombre_hoja in wb.sheetnames:
            hoja = wb[nombre_hoja]
            if norm(nombre_hoja) == "clientes":
                a, o = limpiar_clientes(hoja, clientes, etiquetas)
                print(f"{archivo.relative_to(entrada)} [clientes]: {a} agregados, {o} repetidos")
            elif norm(nombre_hoja) == "productos":
                a, o = limpiar_productos(hoja, productos, existencias, ubicacion, no_publicar)
                print(f"{archivo.relative_to(entrada)} [productos]: {a} agregados, {o} repetidos")

    escribir(salida / "1_etiquetas.xlsx", COLS_ETIQUETAS, etiquetas)
    escribir(salida / "2_clientes.xlsx", COLS_CLIENTES, clientes)
    escribir(salida / "3_productos.xlsx", COLS_PRODUCTOS, productos)
    escribir(salida / "4_existencias.xlsx", COLS_EXISTENCIAS, existencias)
    print(f"\nTOTAL: {len(etiquetas)} etiquetas, {len(clientes)} clientes, {len(productos)} productos, "
          f"{len(existencias)} existencias -> {salida}")


if __name__ == "__main__":
    main()
