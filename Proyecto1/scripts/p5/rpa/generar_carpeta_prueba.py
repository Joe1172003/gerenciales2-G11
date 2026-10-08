"""
Persona 5 - Genera una carpeta de prueba para el bot de UiPath, con la misma
forma que describe el enunciado:

  carpetas  clientes / proveedores / reclamos / registro / productos
  seguidas de un guion y una descripcion, con libros de Excel nombrados igual
  y varias hojas por libro. Solo interesan las hojas "clientes" y "productos".

Incluye a proposito los casos dificiles:
  - las filas y encabezados reales de los archivos de ejemplo del auxiliar
    (Name*, Company Type*, External ID, filas vacias al final, numeros como 1000.0)
  - una hoja "clientes" dentro de un libro de la carpeta proveedores
  - libros sin ninguna hoja util, hojas senuelo y un archivo temporal ~$
  - un cliente y un producto repetidos en dos libros (el bot no debe duplicarlos)
  - un libro dentro de una subcarpeta

Uso:
  python scripts/p5/rpa/generar_carpeta_prueba.py            -> carpeta_prueba_rpa/
  python scripts/p5/rpa/generar_carpeta_prueba.py C:\\RPA\\prueba
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

AQUI = Path(__file__).resolve().parent
EJEMPLOS = AQUI / "ejemplos_auxiliar"
DESTINO = Path(sys.argv[1]) if len(sys.argv) > 1 else AQUI.parents[2] / "carpeta_prueba_rpa"

ENC_CLIENTES = ["Name*", "Company Type*", "Related Company", "Email", "Phone", "Street", "Street2",
                "City", "State", "Zip", "Country", "Tax ID", "Website", "Tags", "Reference", "Notes"]
ENC_PRODUCTOS = ["External ID", "Name", "Product Type", "Internal Reference", "Barcode", "Sales Price",
                 "Cost", "Weight", "Sales Description", "Product Values", "Cantidad a la mano",
                 "Está publicado"]

CLIENTES_GT = [
    ["RPA Abarrotería La Bendición", "Company", None, "compras@labendicion.example.com", "+502 2245 1100",
     "4a. Avenida 12-30", "Zona 1", "Ciudad de Guatemala", "Guatemala", 1001.0, "GT", "1234567-8", None,
     "Mayorista", "CLI-RPA-001", "Cliente de la zona norte"],
    ["RPA María José Pérez", "Person", None, "mariajose.perez@example.com", "+502 5512 3344",
     "Calzada Roosevelt 22-43", None, "Mixco", None, 1057.0, "GT", None, None, None, "CLI-RPA-002", None],
    ["RPA Tienda El Ahorro", "Company", None, "ventas@elahorro.example.com", None, "Blvd. Los Próceres 18",
     None, "San Salvador", None, None, "SV", None, "https://elahorro.example.com", "Mayorista,Frecuente",
     "CLI-RPA-003", None],
    ["RPA Distribuidora Tapachula", "Company", None, "contacto@distap.example.com", "+52 962 555 0101",
     "Av. Central Norte 45", None, "Tapachula", None, 30700.0, "MX", None, None, "Frecuente", "CLI-RPA-004",
     "Pide factura <b>siempre</b>"],
]

PRODUCTOS_EXTRA = [
    ["RPA_PROD_001", "RPA Café molido 454 g", "Goods", "RPA-CAF-454", 7401234500011.0, 38.5, 24.0, 0.454,
     "Café 100 % guatemalteco, tostado medio", None, 40.0, "True"],
    ["RPA_PROD_002", "RPA Servicio de entrega a domicilio", "Service", "RPA-SRV-ENV", None, 25.0, 0.0, None,
     "Entrega en el mismo día dentro del perímetro urbano", None, None, "False"],
]


def leer_filas(nombre):
    """Filas (sin encabezado) de un archivo de ejemplo del auxiliar, tal cual vienen."""
    ws = load_workbook(EJEMPLOS / nombre).active
    filas = list(ws.iter_rows(values_only=True))
    return [list(f) for f in filas[1:]]


def libro(ruta, hojas):
    """hojas: lista de (nombre_hoja, encabezados, filas)."""
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, enc, filas in hojas:
        ws = wb.create_sheet(nombre)
        ws.append(enc)
        for f in filas:
            ws.append(f)
    wb.save(ruta)
    print(f"  {ruta.relative_to(DESTINO)}  ->  hojas: {', '.join(h[0] for h in hojas)}")


def main():
    clientes_aux = leer_filas("clientes-archivo de ejemplo.xlsx")      # 4 con datos + ~215 vacias
    productos_aux = leer_filas("productos - archivo de ejemplo.xlsx")  # 7 con datos + vacias

    reclamos = ("reclamos", ["Fecha", "Cliente", "Motivo", "Estado"],
                [["2026-09-02", "RPA María José Pérez", "Producto vencido", "Cerrado"],
                 ["2026-09-15", "PRUEBA SS 2", "Entrega tardía", "Abierto"]])
    registros = ("registros", ["Fecha", "Usuario", "Acción"],
                 [["2026-09-01", "admin", "Alta de cliente"], ["2026-09-03", "ventas", "Ajuste de precio"]])
    proveedor = ("proveedor", ["Name", "Email", "Country"],
                 [["Distribuidora Centroamericana", "pedidos@dca.example.com", "GT"]])

    print(f"Creando {DESTINO}")
    # 1) clientes: el archivo del auxiliar (con sus filas vacias) mas hojas que no interesan
    libro(DESTINO / "clientes - cartera general" / "clientes - cartera general.xlsx",
          [("clientes", ENC_CLIENTES, clientes_aux), reclamos, registros])
    # 2) clientes en subcarpeta, con un cliente repetido (PRUEBA SS 1) que no debe duplicarse
    libro(DESTINO / "clientes - zona norte" / "2026" / "clientes - zona norte.xlsx",
          [("Hoja1", ["notas"], [["libro de trabajo de la zona norte"]]),
           ("clientes", ENC_CLIENTES, CLIENTES_GT + [clientes_aux[0]])])
    # 3) proveedores: trae una hoja clientes escondida entre otras
    libro(DESTINO / "proveedores - nacionales" / "proveedores - nacionales.xlsx",
          [proveedor, ("clientes", ENC_CLIENTES, [CLIENTES_GT[2]])])
    # 4) reclamos y registro: ninguna hoja util, el bot debe ignorarlos
    libro(DESTINO / "reclamos - enero a septiembre" / "reclamos - enero a septiembre.xlsx", [reclamos])
    libro(DESTINO / "registro - bitacora de ventas" / "registro - bitacora de ventas.xlsx", [registros])
    # 5) productos: el archivo del auxiliar + otro libro con extras y un producto repetido
    libro(DESTINO / "productos - catalogo electronica" / "productos - catalogo electronica.xlsx",
          [("productos", ENC_PRODUCTOS, productos_aux), proveedor])
    libro(DESTINO / "productos - abarrotes y servicios" / "productos - abarrotes y servicios.xlsx",
          [("productos", ENC_PRODUCTOS, PRODUCTOS_EXTRA + [productos_aux[0]]), reclamos])
    # 6) archivo temporal de Excel abierto: el bot debe saltarlo
    tmp = DESTINO / "clientes - cartera general" / "~$clientes - cartera general.xlsx"
    tmp.write_bytes(b"archivo temporal de Excel")
    print(f"  {tmp.relative_to(DESTINO)}  ->  temporal, se debe ignorar")

    print("\nEsperado al procesar: 8 clientes unicos (4 del auxiliar + 4 nuevos; los repetidos no cuentan),"
          " 9 productos unicos (7 del auxiliar + 2 nuevos).")


if __name__ == "__main__":
    main()
