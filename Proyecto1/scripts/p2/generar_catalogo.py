# Genera los dos Excel del catalogo de la tienda, con los encabezados del importador de Odoo
# catalogo_productos.xlsx:   Productos, Importar registros. Tambien es el que se comparte con Persona 4
# catalogo_existencias.xlsx: Operaciones, Inventario fisico, Importar registros. Luego: Aplicar todo
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[2] / 'catalogo'
MODULO = 'catalogo_p2'                  # prefijo del ID externo, para que una segunda carga actualice
SUCURSALES = ('QMGT', 'QMMX', 'QMSV')   # codigos de los almacenes que creo Persona 1
RAIZ = 'All'                            # categoria raiz de Odoo: el importador busca el nombre completo
UBICACION = 'Existencias'               # ubicacion de stock de cada almacen: QMGT/Existencias, etc.

# los mismos colores que usa Persona 1 en los documentos del gestor
CATEGORIAS = {
    'Alimentos': 'B45309',
    'Bebidas': '0369A1',
    'Limpieza': '047857',
    'Hogar': '7C3AED',
}

# categoria, nombre, precio de venta Q, costo Q, peso kg, existencias (QMGT, QMMX, QMSV), descripcion
PRODUCTOS = [
    ('Alimentos', 'Arroz blanco grano largo 1 lb', 7.50, 6.50, 0.454, (180, 120, 90),
     'Arroz de grano largo seleccionado, rinde parejo y no se pega.'),
    ('Alimentos', 'Frijol negro seleccionado 1 lb', 9.25, 8.00, 0.454, (165, 110, 95),
     'Frijol negro limpio y escogido a mano, listo para cocer.'),
    ('Alimentos', 'Aceite vegetal 900 ml', 24.90, 19.90, 0.900, (140, 100, 80),
     'Aceite vegetal puro para freir y sazonar, envase de 900 ml.'),
    ('Alimentos', 'Azucar morena 2.5 kg', 24.50, 20.50, 2.500, (120, 85, 70),
     'Azucar morena de cana guatemalteca en bolsa de 2.5 kg.'),
    ('Alimentos', 'Harina de maiz nixtamalizado 2 kg', 21.75, 17.50, 2.000, (130, 95, 60),
     'Harina nixtamalizada para tortillas y tamalitos, bolsa de 2 kg.'),
    ('Alimentos', 'Cafe molido de altura 454 g', 46.00, 36.00, 0.454, (95, 70, 55),
     'Cafe de altura tostado y molido, cuerpo medio y aroma intenso.'),
    ('Alimentos', 'Pasta espagueti 200 g', 6.75, 5.20, 0.200, (200, 140, 110),
     'Espagueti de semola de trigo, cuece en 8 minutos.'),
    ('Alimentos', 'Atun en agua 140 g', 11.50, 9.00, 0.140, (175, 125, 100),
     'Lomo de atun en agua, sin conservantes, lata de 140 g.'),
    ('Alimentos', 'Leche entera UHT 1 L', 13.90, 11.20, 1.030, (155, 115, 85),
     'Leche entera ultrapasteurizada, no necesita refrigeracion sin abrir.'),
    ('Alimentos', 'Cereal de maiz azucarado 500 g', 32.50, 25.90, 0.500, (110, 80, 65),
     'Hojuelas de maiz con azucar, caja familiar de 500 g.'),

    ('Bebidas', 'Agua pura 600 ml', 3.50, 2.40, 0.600, (240, 180, 150),
     'Agua purificada por osmosis inversa, botella personal de 600 ml.'),
    ('Bebidas', 'Agua pura garrafon 5 galones', 22.00, 16.00, 19.000, (60, 40, 30),
     'Garrafon retornable de 5 galones para dispensador.'),
    ('Bebidas', 'Refresco de cola 2.5 L', 16.50, 12.90, 2.600, (150, 110, 90),
     'Refresco de cola familiar de 2.5 litros, bien carbonatado.'),
    ('Bebidas', 'Jugo de naranja 1 L', 14.75, 11.00, 1.040, (135, 100, 75),
     'Jugo de naranja sin azucar anadida, envase de 1 litro.'),
    ('Bebidas', 'Nectar de mango 1 L', 13.50, 10.20, 1.040, (125, 95, 70),
     'Nectar de mango con pulpa, ideal para el refaccion.'),
    ('Bebidas', 'Bebida hidratante 600 ml', 9.90, 7.20, 0.600, (160, 120, 95),
     'Bebida isotonica con electrolitos, sabor citrico.'),
    ('Bebidas', 'Te helado de limon 500 ml', 8.50, 6.10, 0.500, (145, 105, 80),
     'Te negro frio con limon, endulzado ligero.'),
    ('Bebidas', 'Cerveza lager lata 350 ml', 11.00, 8.40, 0.360, (190, 0, 0),
     'Cerveza lager ligera en lata de 350 ml. Venta solo en Guatemala.'),
    ('Bebidas', 'Horchata en polvo 400 g', 18.90, 14.50, 0.400, (90, 65, 50),
     'Mezcla de horchata morro para preparar en agua o leche.'),
    ('Bebidas', 'Cafe frio listo para tomar 250 ml', 12.50, 9.30, 0.260, (115, 85, 60),
     'Cafe frio con leche, listo para tomar, envase de 250 ml.'),

    ('Limpieza', 'Detergente en polvo 1 kg', 26.90, 21.00, 1.000, (140, 100, 85),
     'Detergente en polvo multiusos con blanqueador, bolsa de 1 kg.'),
    ('Limpieza', 'Cloro 1 galon', 19.50, 14.80, 3.800, (120, 90, 70),
     'Cloro al 5 % para desinfectar pisos y banos, galon.'),
    ('Limpieza', 'Desinfectante multiusos 1 L', 22.75, 17.40, 1.050, (130, 95, 75),
     'Desinfectante de superficies con aroma a lavanda, 1 litro.'),
    ('Limpieza', 'Jabon lavaplatos en crema 425 g', 13.90, 10.50, 0.425, (165, 120, 95),
     'Jabon en crema que corta la grasa y cuida las manos.'),
    ('Limpieza', 'Limpiavidrios 500 ml', 16.50, 12.30, 0.520, (110, 80, 60),
     'Limpiavidrios con atomizador, no deja marcas.'),
    ('Limpieza', 'Suavizante de telas 1 L', 24.90, 19.20, 1.050, (115, 85, 65),
     'Suavizante concentrado con aroma prolongado, 1 litro.'),
    ('Limpieza', 'Papel higienico 4 rollos', 21.50, 16.90, 0.700, (180, 130, 105),
     'Papel higienico doble hoja, paquete de 4 rollos.'),
    ('Limpieza', 'Toallas de cocina 2 rollos', 18.75, 14.20, 0.500, (150, 110, 85),
     'Toallas de papel absorbente para cocina, 2 rollos.'),
    ('Limpieza', 'Bolsas para basura 20 unidades', 15.90, 11.80, 0.650, (160, 115, 90),
     'Bolsas negras resistentes de 20 galones, paquete de 20.'),
    ('Limpieza', 'Esponja multiusos 3 unidades', 9.50, 6.80, 0.090, (175, 125, 100),
     'Esponjas con fibra abrasiva para trastos, paquete de 3.'),

    ('Hogar', 'Escoba de nylon con mango', 32.00, 24.00, 0.600, (45, 35, 25),
     'Escoba de cerda de nylon con mango de metal reforzado.'),
    ('Hogar', 'Trapeador de microfibra', 38.50, 29.00, 0.550, (40, 30, 22),
     'Trapeador de microfibra lavable, atrapa polvo sin dejar pelusa.'),
    ('Hogar', 'Juego de 6 vasos de vidrio', 64.90, 48.00, 2.100, (35, 25, 18),
     'Seis vasos de vidrio templado de 350 ml, aptos para lavavajillas.'),
    ('Hogar', 'Sarten antiadherente 24 cm', 125.00, 92.00, 0.950, (28, 20, 15),
     'Sarten de 24 cm con recubrimiento antiadherente y mango frio.'),
    ('Hogar', 'Olla de aluminio 4 L', 98.50, 72.00, 1.200, (30, 22, 16),
     'Olla de aluminio de 4 litros con tapadera de vidrio.'),
    ('Hogar', 'Juego de 12 ganchos para ropa', 22.90, 16.50, 0.450, (55, 40, 30),
     'Doce ganchos plasticos resistentes para closet.'),
    ('Hogar', 'Caja organizadora plastica 20 L', 74.50, 55.00, 1.100, (32, 24, 18),
     'Caja organizadora de 20 litros con tapadera de cierre.'),
    ('Hogar', 'Bombillo LED 9 W', 27.90, 20.00, 0.080, (70, 50, 40),
     'Bombillo LED de 9 W luz blanca, equivale a 60 W incandescente.'),
    ('Hogar', 'Extension electrica 3 m', 56.00, 41.50, 0.400, (38, 28, 20),
     'Extension de 3 metros con tres tomas y proteccion termica.'),
    ('Hogar', 'Juego de 3 toallas de bano', 149.00, 112.00, 1.500, (25, 18, 12),
     'Tres toallas de algodon peinado: bano, manos y cara.'),
]

# nombre tecnico de cada columna: el importador de Odoo las reconoce en cualquier idioma
COLUMNAS_EXISTENCIAS = ['product_id', 'location_id', 'inventory_quantity']
COLUMNAS = [
    'id', 'name', 'default_code', 'barcode', 'categ_id', 'type', 'is_storable', 'sale_ok',
    'list_price', 'standard_price', 'weight', 'description_sale', 'public_categ_ids',
    'is_published', 'image_1920',
]


def ean13(secuencia):
    """Codigo de barras valido con el prefijo 740 de Guatemala."""
    base = f'7401{secuencia:08d}'
    suma = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(base))
    return base + str((10 - suma % 10) % 10)


def imagen(codigo, nombre, color):
    """Imagen temporal con el nombre del producto. Si conseguimos fotos, se cambia esta columna."""
    texto = quote(f'{codigo}\n{nombre}')
    return f'https://placehold.co/800x800/{color.lower()}/ffffff/png?text={texto}'


productos = []
for n, (categoria, nombre, precio, costo, peso, existencias, descripcion) in enumerate(PRODUCTOS, start=1):
    codigo = f'QMP{n:04d}'
    productos.append({
        'id_externo': f'{MODULO}.{codigo.lower()}',
        'codigo': codigo,
        'nombre': nombre,
        'categoria': categoria,
        'barras': ean13(n),
        'precio': precio,
        'costo': costo,
        'peso': peso,
        'descripcion': descripcion,
        'imagen': imagen(codigo, nombre, CATEGORIAS[categoria]),
        'existencias': dict(zip(SUCURSALES, existencias)),
    })

BASE.mkdir(parents=True, exist_ok=True)

# Excel para el importador de Odoo y para Persona 4
libro = Workbook()
hoja = libro.active
hoja.title = 'productos'
hoja.append(COLUMNAS)
for p in productos:
    hoja.append([
        p['id_externo'], p['nombre'], p['codigo'], p['barras'], f"{RAIZ} / {p['categoria']}", 'consu', True, True,
        p['precio'], p['costo'], p['peso'], p['descripcion'], p['categoria'], True, p['imagen'],
    ])

ANCHOS = {'id': 22, 'name': 36, 'default_code': 14, 'barcode': 16, 'categ_id': 16, 'type': 9,
          'is_storable': 12, 'sale_ok': 10, 'list_price': 12, 'standard_price': 15, 'weight': 9,
          'description_sale': 62, 'public_categ_ids': 18, 'is_published': 13, 'image_1920': 70,
          'product_id': 14, 'location_id': 20, 'inventory_quantity': 20}


def da_formato(hoja, anchos):
    cabecera = PatternFill('solid', fgColor='1D4ED8')
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = cabecera
        celda.alignment = Alignment(horizontal='center')
    for i, columna in enumerate(hoja[1], start=1):
        hoja.column_dimensions[get_column_letter(i)].width = anchos[columna.value]
    hoja.freeze_panes = 'A2'


da_formato(hoja, ANCHOS)
excel = BASE / 'catalogo_productos.xlsx'
libro.save(excel)

# Excel de existencias: una fila por producto y sucursal, para el ajuste de inventario
libro_ex = Workbook()
hoja_ex = libro_ex.active
hoja_ex.title = 'existencias'
hoja_ex.append(COLUMNAS_EXISTENCIAS)
filas = 0
for p in productos:
    for code in SUCURSALES:
        cantidad = p['existencias'][code]
        if not cantidad:
            continue
        hoja_ex.append([p['codigo'], f'{code}/{UBICACION}', cantidad])
        filas += 1
da_formato(hoja_ex, ANCHOS)
excel_ex = BASE / 'catalogo_existencias.xlsx'
libro_ex.save(excel_ex)

print('OK', excel.name, len(productos), 'productos')
for categoria in CATEGORIAS:
    de_la_categoria = [p for p in productos if p['categoria'] == categoria]
    print('OK', categoria, len(de_la_categoria), 'productos |',
          de_la_categoria[0]['codigo'], 'a', de_la_categoria[-1]['codigo'])
for code in SUCURSALES:
    print('OK existencias', code, sum(p['existencias'][code] for p in productos), 'unidades')
print('OK', excel_ex.name, filas, 'filas de ajuste de inventario')
